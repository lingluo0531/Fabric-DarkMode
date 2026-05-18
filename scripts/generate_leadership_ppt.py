from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt


OUTPUT_PATH = Path(__file__).resolve().parents[1] / "slides" / "dark-mode-private-preview-leadership-review-v2.pptx"


def set_run_font(run, size, bold=False, color=(0, 0, 0), name="Aptos"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)


def add_textbox(slide, left, top, width, height, text, size, color, bold=False, margin=0.06):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run_font(run, size=size, bold=bold, color=color)
    return box


def add_panel(slide, left, top, width, height, fill_rgb, line_rgb=(220, 226, 232)):
    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(*fill_rgb)
    panel.line.color.rgb = RGBColor(*line_rgb)
    panel.line.width = Pt(1)
    return panel


def add_bullets(slide, left, top, width, height, items, font_size=12, color=(22, 33, 43)):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.04)
    frame.clear()
    for i, text in enumerate(items):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.level = 0
        p.bullet = True
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = text
        set_run_font(run, size=font_size, color=color)
    return box


def add_banner(slide, label):
    top_band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.34))
    top_band.fill.solid()
    top_band.fill.fore_color.rgb = RGBColor(21, 35, 46)
    top_band.line.fill.background()

    right_band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(12.95), Inches(0), Inches(0.383), Inches(7.5))
    right_band.fill.solid()
    right_band.fill.fore_color.rgb = RGBColor(15, 108, 189)
    right_band.line.fill.background()

    add_textbox(slide, Inches(0.45), Inches(0.16), Inches(4.2), Inches(0.22), label, 10, (255, 255, 255), bold=True)


def add_slide_one(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(244, 239, 230)

    add_banner(slide, "LEADERSHIP REVIEW | FABRIC DARK MODE")

    add_textbox(
        slide,
        Inches(0.5),
        Inches(0.62),
        Inches(8.0),
        Inches(0.95),
        "Current Status: Private Preview is on track with a clear and controlled scope.",
        25,
        (22, 33, 43),
        bold=True,
    )
    add_textbox(
        slide,
        Inches(0.52),
        Inches(1.48),
        Inches(8.0),
        Inches(0.72),
        "Dark mode now has stable adoption signals in private preview scope, with safe fallback behavior for surfaces not yet onboarded.",
        12,
        (83, 97, 112),
    )

    add_panel(slide, Inches(9.05), Inches(0.62), Inches(3.63), Inches(1.5), (21, 35, 46), line_rgb=(21, 35, 46))
    add_textbox(slide, Inches(9.24), Inches(0.83), Inches(3.2), Inches(0.2), "PRIVATE PREVIEW TARGET", 9.5, (202, 217, 227), bold=True)
    add_textbox(slide, Inches(9.2), Inches(1.03), Inches(3.2), Inches(0.42), "May 29, 2026", 24, (255, 255, 255), bold=True)
    add_textbox(slide, Inches(9.2), Inches(1.43), Inches(3.2), Inches(0.38), "Status: On track for planned private preview milestone.", 10.5, (221, 231, 238))

    add_panel(slide, Inches(0.45), Inches(2.28), Inches(3.45), Inches(4.35), (255, 252, 247))
    add_textbox(slide, Inches(0.62), Inches(2.43), Inches(2.8), Inches(0.22), "PRIVATE PREVIEW SCOPE", 10, (83, 97, 112), bold=True)
    scope_items = [
        "Fabric Shell",
        "Workloads Hub L1",
        "OneLake Category",
        "Lakehouse",
        "Variable Library",
        "Notebook",
        "Real-Time Hub",
        "Kusto",
        "Warehouse and Azure SQL DW",
        "Data Science",
    ]
    add_bullets(slide, Inches(0.6), Inches(2.72), Inches(3.1), Inches(3.7), scope_items, font_size=11)

    add_panel(slide, Inches(4.05), Inches(2.28), Inches(4.8), Inches(2.03), (255, 252, 247))
    add_textbox(slide, Inches(4.22), Inches(2.43), Inches(4.3), Inches(0.22), "CURRENT ADOPTION STATUS", 10, (83, 97, 112), bold=True)
    status_items = [
        "Strong signal: bug bash completed, no sev1/sev2 blockers.",
        "Activation path works: users can turn on dark mode and persist preference.",
        "Scope is intentionally controlled for quality during preview.",
    ]
    add_bullets(slide, Inches(4.18), Inches(2.73), Inches(4.45), Inches(1.43), status_items, font_size=10.5)

    add_panel(slide, Inches(4.05), Inches(4.5), Inches(4.8), Inches(2.13), (255, 252, 247))
    add_textbox(slide, Inches(4.22), Inches(4.65), Inches(4.3), Inches(0.22), "WHY THIS MATTERS", 10, (83, 97, 112), bold=True)
    value_items = [
        "Reduces eye strain in long creator sessions.",
        "Improves platform modernity and premium perception.",
        "Builds confidence with consistent theme behavior in core flows.",
    ]
    add_bullets(slide, Inches(4.18), Inches(4.95), Inches(4.45), Inches(1.46), value_items, font_size=10.5)

    add_panel(slide, Inches(9.05), Inches(2.28), Inches(3.63), Inches(4.35), (255, 252, 247))
    add_textbox(slide, Inches(9.24), Inches(2.43), Inches(3.15), Inches(0.22), "PRIVATE PREVIEW STORYLINE", 10, (83, 97, 112), bold=True)
    story_items = [
        "Enable dark mode from Settings.",
        "Refresh and verify preference persistence.",
        "Navigate key shell surfaces and supported artifacts.",
        "Show safe fallback where onboarding is not yet complete.",
    ]
    add_bullets(slide, Inches(9.2), Inches(2.72), Inches(3.2), Inches(1.6), story_items, font_size=10)

    add_textbox(
        slide,
        Inches(9.2),
        Inches(4.5),
        Inches(3.2),
        Inches(1.75),
        "Private preview validates quality and readiness, while full platform consistency depends on the next onboarding wave after preview.",
        10.2,
        (83, 97, 112),
    )


def add_table(slide, left, top, col_widths, headers, rows, row_height=0.36, font_size=8.7):
    h = Inches(0.34)
    x = left
    for i, head in enumerate(headers):
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, top, col_widths[i], h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(207, 220, 234)
        shape.line.color.rgb = RGBColor(150, 160, 170)
        shape.line.width = Pt(0.5)
        add_textbox(slide, x + Inches(0.05), top + Inches(0.05), col_widths[i] - Inches(0.08), Inches(0.22), head, 9, (22, 33, 43), bold=True)
        x += col_widths[i]

    y = top + h
    row_h = Inches(row_height)
    for r, row in enumerate(rows):
        row_color = (255, 255, 255) if r % 2 == 0 else (247, 249, 250)
        x = left
        for i, val in enumerate(row):
            cell = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, col_widths[i], row_h)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(*row_color)
            cell.line.color.rgb = RGBColor(215, 220, 228)
            cell.line.width = Pt(0.5)
            add_textbox(slide, x + Inches(0.05), y + Inches(0.04), col_widths[i] - Inches(0.08), row_h - Inches(0.06), str(val), font_size, (22, 33, 43), bold=False)
            x += col_widths[i]
        y += row_h


def add_slide_two(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(244, 239, 230)

    add_banner(slide, "LEADERSHIP REVIEW | FULL ADOPTION PLAN")

    add_textbox(
        slide,
        Inches(0.5),
        Inches(0.62),
        Inches(8.65),
        Inches(0.9),
        "Post-Preview Onboarding: close the gap and achieve full adoption by Sep 30, 2026.",
        24,
        (22, 33, 43),
        bold=True,
    )
    add_textbox(
        slide,
        Inches(0.52),
        Inches(1.43),
        Inches(8.65),
        Inches(0.65),
        "Execution objective: condensed onboarding motion across remaining teams to minimize the mixed light/dark UI window.",
        12,
        (83, 97, 112),
    )

    add_panel(slide, Inches(9.05), Inches(0.62), Inches(3.63), Inches(1.5), (21, 35, 46), line_rgb=(21, 35, 46))
    add_textbox(slide, Inches(9.24), Inches(0.83), Inches(3.2), Inches(0.2), "FULL ADOPTION TARGET", 9.5, (202, 217, 227), bold=True)
    add_textbox(slide, Inches(9.2), Inches(1.03), Inches(3.2), Inches(0.42), "Sep 30, 2026", 24, (255, 255, 255), bold=True)
    add_textbox(slide, Inches(9.2), Inches(1.43), Inches(3.2), Inches(0.38), "Goal: pull deferred teams into RB semester delivery.", 10.5, (221, 231, 238))

    add_panel(slide, Inches(0.45), Inches(2.2), Inches(8.45), Inches(2.95), (255, 252, 247))
    add_textbox(slide, Inches(0.62), Inches(2.35), Inches(6.8), Inches(0.22), "POST-PREVIEW ONBOARDING ITEMS AND TIMELINE (RB + POST-RB)", 10, (83, 97, 112), bold=True)
    headers = ["Team", "Scenario", "Solution", "Cost", "Prior Plan"]
    rows = [
        ["OneLake Catalog - Govern", "RB", "Onboard in RB train", "~1-2 weeks", "RB"],
        ["Deployment Pipelines", "RB", "Onboard in RB train", "~1-2 weeks", "RB"],
        ["Event Stream", "RB", "Onboard in RB train", "~1-2 weeks", "RB"],
        ["Azure Data Factory", "RB", "Onboard in RB train", "~1-2 weeks", "RB"],
        ["Spark Job Definition", "RB", "Onboard in RB train", "~1-2 weeks", "RB"],
        ["Power BI", "PBIClient", "Full validation + AI legacy token fixes", "Variable", "Post-RB"],
        ["Dataflow", "React + Fluent + 1JS", "Upgrade + ExtensionThemeProvider", "~1 day", "Post-RB"],
        ["Data Activator", "TBD", "Confirm scenario and run standard playbook", "TBD", "RB or Post-RB"],
        ["Functions Hub", "TBD", "Confirm scenario and run standard playbook", "TBD", "Post-RB"],
    ]
    col_widths = [Inches(1.45), Inches(1.45), Inches(2.9), Inches(0.95), Inches(1.3)]
    add_table(slide, Inches(0.6), Inches(2.67), col_widths, headers, rows, row_height=0.26, font_size=7.8)

    add_panel(slide, Inches(9.05), Inches(2.2), Inches(3.63), Inches(2.95), (255, 252, 247))
    add_textbox(slide, Inches(9.24), Inches(2.35), Inches(3.2), Inches(0.22), "ACCELERATION REQUIRED BEFORE 9/30", 10, (83, 97, 112), bold=True)
    gap_headers = ["Team", "Owner", "Current Plan"]
    gap_rows = [
        ["Data Activator", "James Hutton", "RB or Post-RB"],
        ["Functions Hub", "Sunitha Muthukrishna", "Post-RB"],
        ["Power BI", "Kay Unkroth", "Post-RB"],
        ["Dataflow", "Miguel Escobar", "Post-RB"],
    ]
    gap_col_widths = [Inches(1.18), Inches(1.35), Inches(0.9)]
    add_table(slide, Inches(9.18), Inches(2.67), gap_col_widths, gap_headers, gap_rows, row_height=0.36, font_size=7.8)
    add_textbox(slide, Inches(9.18), Inches(4.48), Inches(3.3), Inches(0.5), "Gap highlight: these four teams must be pulled into the 9/30 plan to reduce the mixed light/dark UI window.", 9, (180, 35, 24), bold=True)

    add_panel(slide, Inches(0.45), Inches(5.32), Inches(12.23), Inches(1.95), (255, 252, 247))
    add_textbox(slide, Inches(0.62), Inches(5.47), Inches(5.8), Inches(0.22), "FULL ADOPTION PLAN", 10, (83, 97, 112), bold=True)
    plan_items = [
        "Step 1 (Week 1): Sync with all four teams and lock scenario mapping. Known: Power BI -> Scenario 1, Dataflow -> Scenario 2.",
        "Step 2 (Week 1-2): Get commitment from each team on completion by 9/30, with named blockers and escalation owner.",
        "Step 3 (Week 2+): Run regular checkpoint review and keep a single cross-team tracker to avoid drift.",
        "Step 4 (By 9/30): Exit gate for full adoption: implementation done, key flow validation complete, no blocking inconsistency defects.",
    ]
    add_bullets(slide, Inches(0.6), Inches(5.75), Inches(11.95), Inches(1.33), plan_items, font_size=10.2)


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_slide_one(prs)
    add_slide_two(prs)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    out = build_deck()
    print(out)
